import os
from datetime import datetime
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.models.contract_spec import ContractSpec

# Formatting constants
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
GROUP_FILL = PatternFill("solid", fgColor="D6E4F0")
GROUP_FONT = Font(name="Calibri", size=10, bold=True)
DATA_FONT = Font(name="Calibri", size=10)
DATA_FONT_CN = Font(name="Microsoft YaHei", size=10)
THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
)
ALT_ROW_FILL = PatternFill("solid", fgColor="F2F2F2")
HYPERLINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")

# Column definitions for each sheet type
BASIC_COLUMNS = [
    ("exchange_code", "Exchange", 10),
    ("exchange_sub", "Sub-Exchange", 14),
    ("product_name_en", "Product (EN)", 30),
    ("product_name_cn", "Product (CN)", 20),
    ("product_group", "Group", 16),
    ("ticker", "Ticker", 10),
    ("bloomberg_code", "Bloomberg", 14),
    ("contract_size", "Contract Size", 25),
    ("quote_unit", "Quote Unit", 20),
    ("quote_currency", "Currency", 10),
    ("tick_size", "Tick Size", 25),
    ("tick_size_numeric", "Tick (Numeric)", 14),
    ("tick_value", "Tick Value", 12),
    ("contract_months", "Contract Months", 30),
    ("listed_contracts", "Listed Contracts", 18),
    ("trading_hours_electronic", "Electronic Hours", 30),
    ("trading_hours_regular", "Regular Hours", 25),
    ("settlement_method", "Settlement", 15),
]

TAS_COLUMNS = [
    ("exchange_code", "Exchange", 10),
    ("exchange_sub", "Sub-Exchange", 14),
    ("product_name_en", "Product", 30),
    ("ticker", "Ticker", 10),
    ("tas_eligible", "TAS", 6),
    ("tas_tick_range", "TAS Tick Range", 25),
    ("tas_eligible_months", "TAS Months", 30),
    ("tas_trading_hours", "TAS Hours", 25),
    ("tam_eligible", "TAM", 6),
    ("tam_rules", "TAM Rules", 30),
    ("btic_eligible", "BTIC", 6),
    ("btic_rules", "BTIC Rules", 25),
    ("block_trade_minimum", "Block Min", 15),
    ("efp_eligible", "EFP", 6),
]

RISK_COLUMNS = [
    ("exchange_code", "Exchange", 10),
    ("exchange_sub", "Sub-Exchange", 14),
    ("product_name_en", "Product", 30),
    ("ticker", "Ticker", 10),
    ("price_limit_daily", "Daily Price Limit", 30),
    ("price_limit_type", "Limit Type", 18),
    ("initial_margin", "Initial Margin", 18),
    ("maintenance_margin", "Maint. Margin", 18),
    ("margin_currency", "Margin Ccy", 10),
    ("position_limit_spot", "Pos Limit (Spot)", 18),
    ("position_limit_single", "Pos Limit (Single)", 18),
    ("position_limit_all", "Pos Limit (All)", 18),
    ("large_trader_reporting", "Reportable Level", 18),
]

DELIVERY_COLUMNS = [
    ("exchange_code", "Exchange", 10),
    ("exchange_sub", "Sub-Exchange", 14),
    ("product_name_en", "Product", 30),
    ("ticker", "Ticker", 10),
    ("delivery_method", "Delivery Method", 16),
    ("last_trading_day", "Last Trading Day", 35),
    ("first_notice_day", "First Notice Day", 30),
    ("first_delivery_day", "First Delivery", 25),
    ("last_delivery_day", "Last Delivery", 25),
    ("delivery_period", "Delivery Period", 25),
    ("deliverable_grades", "Deliverable Grades", 40),
    ("delivery_points", "Delivery Points", 30),
    ("final_settlement_price", "Final Settlement", 35),
]

ALL_COLUMNS = BASIC_COLUMNS + [
    c for c in TAS_COLUMNS if c[0] not in {b[0] for b in BASIC_COLUMNS}
] + [
    c for c in RISK_COLUMNS if c[0] not in {b[0] for b in BASIC_COLUMNS} and c[0] not in {t[0] for t in TAS_COLUMNS}
] + [
    c for c in DELIVERY_COLUMNS if c[0] not in {b[0] for b in BASIC_COLUMNS}
    and c[0] not in {t[0] for t in TAS_COLUMNS} and c[0] not in {r[0] for r in RISK_COLUMNS}
]


def export_to_excel(
    contracts: List[ContractSpec],
    output_dir: str = "data/output",
    filename: Optional[str] = None,
) -> str:
    """Export contract specs to a professionally formatted Excel workbook."""
    os.makedirs(output_dir, exist_ok=True)
    if filename is None:
        filename = f"Global_Futures_Specs_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = Workbook()

    # Sheet 1: Summary
    _build_summary_sheet(wb.active, contracts)
    wb.active.title = "Summary"

    # Sheet 2: All Contracts
    ws_all = wb.create_sheet("All Contracts")
    _build_data_sheet(ws_all, contracts, ALL_COLUMNS)

    # Sheet 3+: Per-exchange sheets
    exchanges = sorted(set(c.exchange_code for c in contracts))
    for exch in exchanges:
        exch_contracts = [c for c in contracts if c.exchange_code == exch]
        ws = wb.create_sheet(exch)
        _build_data_sheet(ws, exch_contracts, BASIC_COLUMNS)

    # TAS-TAM Rules sheet
    tas_contracts = [c for c in contracts if c.tas_eligible or c.tam_eligible or c.btic_eligible]
    if tas_contracts:
        ws_tas = wb.create_sheet("TAS-TAM Rules")
        _build_data_sheet(ws_tas, tas_contracts, TAS_COLUMNS)

    # Risk Controls sheet
    risk_contracts = [c for c in contracts if any([
        c.price_limit_daily, c.initial_margin, c.position_limit_spot, c.position_limit_all
    ])]
    if risk_contracts:
        ws_risk = wb.create_sheet("Risk Controls")
        _build_data_sheet(ws_risk, risk_contracts, RISK_COLUMNS)

    # Delivery Specs sheet
    delivery_contracts = [c for c in contracts if c.delivery_method is not None]
    if delivery_contracts:
        ws_del = wb.create_sheet("Delivery Specs")
        _build_data_sheet(ws_del, delivery_contracts, DELIVERY_COLUMNS)

    # Data Quality sheet
    ws_quality = wb.create_sheet("Data Quality")
    _build_quality_sheet(ws_quality, contracts)

    wb.save(filepath)
    return filepath


def _build_summary_sheet(ws: Worksheet, contracts: List[ContractSpec]):
    """Build the summary dashboard sheet."""
    ws.sheet_properties.tabColor = "1F4E79"

    # Title
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = "Global Futures Contract Specifications"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:F2")
    ws["A2"].value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="666666")
    ws["A2"].alignment = Alignment(horizontal="center")

    # Summary table
    headers = ["Exchange", "Products", "Avg Completeness %", "TAS Eligible", "Region", "Last Updated"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    exchanges = sorted(set(c.exchange_code for c in contracts))
    region_map = _get_region_map()

    for row_idx, exch in enumerate(exchanges, 5):
        exch_contracts = [c for c in contracts if c.exchange_code == exch]
        avg_completeness = sum(c.completeness_score() for c in exch_contracts) / len(exch_contracts)
        tas_count = sum(1 for c in exch_contracts if c.tas_eligible)
        region = region_map.get(exch, "Other")
        last_updated = max((c.last_updated or "" for c in exch_contracts), default="N/A")

        values = [exch, len(exch_contracts), f"{avg_completeness:.1f}%", tas_count, region, last_updated]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if row_idx % 2 == 0:
                cell.fill = ALT_ROW_FILL

    # Total row
    total_row = 5 + len(exchanges)
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(name="Calibri", size=10, bold=True)
    ws.cell(row=total_row, column=2, value=len(contracts)).font = Font(name="Calibri", size=10, bold=True)

    # Column widths
    widths = [14, 12, 20, 14, 12, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A5"


def _build_data_sheet(ws: Worksheet, contracts: List[ContractSpec], columns: list):
    """Build a data sheet with headers, data rows, alternating colors, and autofilter."""
    # Headers
    for col_idx, (field, label, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data rows
    for row_idx, spec in enumerate(contracts, 2):
        data = spec.model_dump()
        for col_idx, (field, _, _) in enumerate(columns, 1):
            val = data.get(field)
            if isinstance(val, bool):
                val = "Yes" if val else "No"
            elif val is not None:
                val = str(val)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if row_idx % 2 == 0:
                cell.fill = ALT_ROW_FILL

    # AutoFilter
    if contracts:
        last_col = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A1:{last_col}{len(contracts) + 1}"

    # Freeze header and first 3 columns
    ws.freeze_panes = "D2"


def _build_quality_sheet(ws: Worksheet, contracts: List[ContractSpec]):
    """Build the data quality audit sheet."""
    headers = ["Exchange", "Product", "Ticker", "Completeness %", "Quality", "Source Type", "Source URL", "Updated"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for row_idx, spec in enumerate(contracts, 2):
        values = [
            spec.exchange_code,
            spec.product_name_en,
            spec.ticker,
            f"{spec.completeness_score():.1f}%",
            spec.data_quality or "unknown",
            spec.source_type or "N/A",
            spec.source_url or "N/A",
            spec.last_updated or "N/A",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if row_idx % 2 == 0:
                cell.fill = ALT_ROW_FILL
            # Hyperlink for source URL
            if col == 7 and val != "N/A" and val.startswith("http"):
                cell.font = HYPERLINK_FONT
                cell.hyperlink = val

    widths = [12, 30, 10, 16, 14, 12, 50, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _get_region_map() -> dict:
    return {
        "CME": "North America", "CBOT": "North America", "NYMEX": "North America", "COMEX": "North America",
        "ICE_US": "North America", "CBOE": "North America",
        "ICE_EU": "Europe", "EUREX": "Europe", "LME": "Europe",
        "SGX": "Asia-Pacific", "HKEX": "Asia-Pacific", "JPX": "Asia-Pacific",
        "KRX": "Asia-Pacific", "TAIFEX": "Asia-Pacific", "BMD": "Asia-Pacific",
        "NSE": "Asia-Pacific", "ASX": "Asia-Pacific",
        "SHFE": "China", "DCE": "China", "CZCE": "China",
        "CFFEX": "China", "INE": "China", "GFEX": "China",
        "B3": "Americas", "MOEX": "Europe",
    }
